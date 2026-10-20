"""Build the P12 report from page-delimited Markdown using bundled python-docx."""
import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def inline(paragraph, text):
    for i, part in enumerate(re.split(r'`([^`]+)`', text)):
        run = paragraph.add_run(part)
        if i % 2:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)


def build(source, output, assets):
    pages = source.read_text(encoding='utf-8').split('<!-- page -->')
    assert 25 <= len(pages) <= 35, len(pages)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
    section.left_margin, section.right_margin = Cm(2.6), Cm(2.2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
    section.different_first_page_header_footer = True
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal = document.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 22), ('Heading 1', 17), ('Heading 2', 14)]:
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.bold = True
    document.styles['Caption'].font.size = Pt(10)
    document.styles['Caption'].font.italic = True
    header = section.header.paragraphs[0]
    header.text = 'ĐỒ ÁN CHUYÊN NGÀNH  •  FACE MASK DETECTION'
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('P12 để review  •  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    toc = [(re.search(r'^# (.+)', page.strip(), re.M).group(1), n+1) for n, page in enumerate(pages)]
    figure_names = {'image', 'video', 'webcam', 'architecture', 'sequence'}
    for index, page in enumerate(pages):
        if index:
            document.add_page_break()
        lines = page.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1; continue
            if line == '{{toc}}':
                for title, number in toc[3:]:
                    if not title.startswith(('Chương', 'Kết luận', 'Tài liệu', 'Phụ lục')):
                        continue
                    p = document.add_paragraph()
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT)
                    r=p.add_run(f'{title}\t{number}');r.font.size=Pt(10)
            elif line.startswith('{{'):
                name = line.strip('{}')
                assert name in figure_names
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Screenshots are full app frames, scaled without cropping.
                p.add_run().add_picture(str(assets / f'{name}.png'), width=Cm(9 if name=='video' else 11.5 if name=='webcam' else 15.5))
                p.paragraph_format.keep_with_next = True
            elif line.startswith('|'):
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    cells=[s.strip() for s in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r':?-+:?', s.replace(' ', '')) for s in cells):
                        rows.append(cells)
                    i += 1
                table = document.add_table(rows=0, cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                widths = ([.25,.75] if rows[0][0]=='Thuật ngữ' else
                          [.31,.13,.19,.19,.18] if rows[0][0]=='Lớp test' else
                          [.30,.14,.14,.14,.14,.14] if rows[0][0]=='Split' and len(rows[0])==6 else
                          [1/len(rows[0])]*len(rows[0]))
                for column, width in zip(table.columns,widths):column.width=Cm(16.2*width)
                for rno, row in enumerate(rows):
                    cells=table.add_row().cells
                    for cno, text in enumerate(row):
                        cell=cells[cno];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                        p.paragraph_format.line_spacing=1.05
                        inline(p,text)
                        for run in p.runs:run.font.size=Pt(9.5);run.bold=(rno==0)
                        props=cell._tc.get_or_add_tcPr()
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DFE8EF' if rno==0 else ('F5F7F9' if rno%2==0 else 'FFFFFF'));props.append(shade)
                        borders=OxmlElement('w:tcBorders')
                        for edge in ['top','left','bottom','right']:
                            b=OxmlElement(f'w:{edge}');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                        props.append(borders)
                    if rno==0:
                        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                document.add_paragraph().paragraph_format.space_after=Pt(0)
                continue
            elif line.startswith('```'):
                i+=1
                while i<len(lines) and not lines[i].startswith('```'):
                    p=document.add_paragraph(lines[i]);p.paragraph_format.space_after=Pt(1)
                    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
                    i+=1
            elif line.startswith('# '):
                p=document.add_paragraph(line[2:], 'Title' if index==0 else 'Heading 1')
                if index==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            elif line.startswith('## '):
                document.add_paragraph(line[3:], 'Heading 2')
            else:
                style='Caption' if line.startswith('Hình ') else None
                p=document.add_paragraph(style=style);inline(p,line)
                if index==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                elif not style:p.alignment=WD_ALIGN_PARAGRAPH.LEFT if 'SHA-256:' in line or index==24 else WD_ALIGN_PARAGRAPH.JUSTIFY
            i+=1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title='Xây dựng ứng dụng web phát hiện trạng thái đeo khẩu trang'
    document.core_properties.subject='Báo cáo đồ án P12 để review, evidence P4–P11'
    document.core_properties.author='Nhóm 1'
    document.save(output)
    print(f'Built {len(pages)} planned pages: {output}')


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--assets',type=Path,required=True)
    args=parser.parse_args()
    build(Path(__file__).with_name('bao-cao.md'),args.output,args.assets)
