"""Assert source fidelity, refreshed lists, original metrics and unchanged P12 outputs."""
import argparse
import hashlib
import json
import posixpath
import re
import unicodedata
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
from build_report import REFERENCE_SHA


def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def compact(s):return re.sub(r'\s+','',unicodedata.normalize('NFC',s)).casefold()
def vi(v,n):return f'{v:,.{n}f}'.translate(str.maketrans({',':'.','.':','}))
def roman(n):
    result=''
    for value,letter in [(10,'x'),(9,'ix'),(5,'v'),(4,'iv'),(1,'i')]:
        while n>=value:result+=letter;n-=value
    return result


def verify(reference,output,work,evidence):
    assert digest(reference)==REFERENCE_SHA
    d=Document(output/'bao-cao-do-an.docx');pdf=PdfReader(output/'bao-cao-do-an.pdf')
    receipt=load(work/'word-final.json');authored=load(work/'authoring-receipt.json')
    assert receipt['read_only'] and receipt['toc_count']==3
    assert len(pdf.pages)==receipt['pages'] and len(pdf.pages)>0
    assert len(d.sections)==3
    for section in d.sections[1:]:
        assert abs(section.top_margin.cm-3)<.02
        for name in ['bottom_margin','left_margin','right_margin']:assert abs(getattr(section,name).cm-2.54)<.02
    normal=d.styles['Normal']
    defaults=d.styles.element.xpath('./w:docDefaults')[0]
    assert (normal.font.name or defaults.xpath('.//*[local-name()="rFonts"]')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii'))=='Times New Roman'
    assert (normal.font.size.pt if normal.font.size else int(defaults.xpath('.//*[local-name()="sz"]')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))/2)==13
    assert (normal.paragraph_format.line_spacing or int(defaults.xpath('.//*[local-name()="spacing"]')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}line'))/240)==1.5
    texts=[p.extract_text() for p in pdf.pages];all_text=compact('\n'.join(texts))
    paragraphs=[p.text for p in d.paragraphs]
    for phrase in ['Ngoại lệ: loại/dung lượng sai', 'Không vẽ bbox của frame cũ', 'Ngoại lệ: API chậm', 'Ngoại lệ: từ chối quyền']:
        assert any(p.startswith(phrase) and p.rstrip().endswith('.') for p in paragraphs),('fragmented paragraph',phrase)
    for p in d.paragraphs:
        if p.text.startswith('[7] TorchVision.'):
            assert any(r.font.size.pt==8 and r.text.startswith('https://docs.pytorch.org/') for r in p.runs)
            assert compact('https://docs.pytorch.org/vision/stable/models/generated/torchvision.models.detection.ssdlite320_mobilenet_v3_large.html') in all_text
    for bad in ['UniLend','Uni-Lend','Nguyễn Xuân Hoàng','Phạm Văn Hà','Dư Quốc An','20261IT6040001']:
        assert compact(bad) not in all_text,bad
    assert 'LỜI CẢM ƠN' not in texts[0] and 'LỜI CẢM ƠN' in texts[1]
    for value in ['ThS. Vũ Thị Dương','2026IT6150003','Phùng Đức Anh','Nguyễn Đình Luân Kiệt','Bùi Nguyên Hải','Bùi Ngọc Minh Hoàng']:assert compact(value) in compact(texts[0])
    toc=[p.text for p in d.paragraphs if p.style.name.lower().startswith('toc')]
    chapter_entries=[e for e in receipt['entries'] if e['style']=='Heading 1' and e['text'].startswith('CHƯƠNG ')]
    assert len(chapter_entries)==5
    matched=[]
    for entry in receipt['entries']:
        title=compact(entry['text']);physical=entry['physical_page'];page=entry['section_page']
        assert title in compact(texts[physical-1]),('PDF location',entry)
        expected=roman(page) if physical<chapter_entries[0]['physical_page'] else str(page)
        matches=[t for t in toc if '\t' in t and compact(t.rsplit('\t',1)[0])==title]
        assert matches and any(t.rsplit('\t',1)[1].strip()==expected for t in matches),('stale list entry',entry,matches)
        matched.append({'text':entry['text'],'physical_page':physical,'printed_page':expected})
    for stage in ['test','validation']:
        report=load(evidence/f'p4-evaluation/{stage}/report.json')
        for metric in report['operating']['macro'].values():assert compact(vi(metric,4)) in all_text
        for key in ['ap50','map50_95']:assert compact(vi(report['ap'][key],4)) in all_text
        if stage=='test':
            for values in report['operating']['per_class'].values():
                for key in ['precision','recall','f1']:assert compact(vi(values[key],4)) in all_text
    api=load(evidence/'p11-internal-regression/real-model-run-3/api-summary.json')
    browser=load(evidence/'p11-internal-regression/real-model-run-3/browser-summary.json')
    for kind in ['cold','warm']:
        for key in ['p50_ms','p95_ms']:assert compact(vi(api[kind][key],2)) in all_text
    for kind in ['video','webcam']:
        for key in ['displayed_frames_per_s','duration_s','p95_ms']:assert compact(vi(browser[kind][key],2)) in all_text
    for important in ['0,001','0,15','uncertain','CLOSED','562','356,83','102,09','7,96','2/49','1/12']:
        assert compact(important) in all_text,important
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(output/'bao-cao-do-an.docx') as z:
        names=set(z.namelist())
        assert all(hashlib.sha256(z.read(r['part'])).hexdigest()==r['sha256'] for r in authored['preserved_parts'])
        root=E.fromstring(z.read('word/document.xml'))
        instructions=root.xpath('//*[local-name()="instrText"]/text()')
        assert sum(t.strip().startswith('TOC ') for t in instructions)==3
        assert sum(t.strip().startswith('SEQ ') for t in instructions)==33
        for n in names:
            if n.endswith('.rels'):
                for rel in E.fromstring(z.read(n)):
                    if rel.get('TargetMode')=='External':continue
                    resolved=posixpath.normpath(posixpath.join(posixpath.dirname(posixpath.dirname(n)),rel.get('Target'))).lstrip('/')
                    assert resolved in names,(n,resolved)
        # python-docx load above additionally resolves restored font content types.
    old_manifest=load(evidence/'p12-report-slides/evidence-manifest.json')
    for item in old_manifest['records']:
        if 'deliverable' in item:assert digest(evidence/'p12-report-slides'/item['deliverable'])==item['sha256']
    assert len([e for e in receipt['entries'] if e['style']=='FigureCaption'])==6
    assert len([e for e in receipt['entries'] if e['style']=='TableCaption'])==27
    result={'status':'PASS','pages':len(pdf.pages),'figures':6,'tables':27,'toc_entries_checked':len(matched),
        'word_version':receipt['word_version'],'reference_sha256':REFERENCE_SHA,
        'preserved_parts_checked':len(authored['preserved_parts']),'old_p12_unchanged':True,
        'independent_review':False,'deployment_gate':'CLOSED','entries':matched,
        'deliverables':[{'name':name,'sha256':digest(output/name),'bytes':(output/name).stat().st_size} for name in ['bao-cao-do-an.docx','bao-cao-do-an.pdf']]}
    (output/'manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'PASS: {len(pdf.pages)} pages, 6 figures, 27 tables, {len(matched)} refreshed list entries, original metrics and P12 hashes.')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--reference',type=Path,required=True);p.add_argument('--output',type=Path,required=True);p.add_argument('--work',type=Path,required=True);p.add_argument('--evidence-root',type=Path,default=Path('D:/Codex/outputs/face-mask-detection'))
    a=p.parse_args();verify(a.reference,a.output,a.work,a.evidence_root)
