"""Adapt the supplied DOCX package, preserving its cover frame and opaque parts."""
import argparse
import copy
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE

REFERENCE_SHA='7d8b433eefcb6c7d53484370d4e775b88daa60337595afa6ee13a72b7fddde5a'


def xml(tag, **attrs):
    node=OxmlElement('w:'+tag)
    for key,value in attrs.items():node.set(qn('w:'+key),str(value))
    return node


def field(p, instruction, cached=''):
    for node in [xml('fldChar',fldCharType='begin'),xml('instrText'),xml('fldChar',fldCharType='separate')]:
        if node.tag==qn('w:instrText'):node.text=' '+instruction+' '
        p.add_run()._r.append(node)
    p.add_run(cached)
    p.add_run()._r.append(xml('fldChar',fldCharType='end'))


def inline(p,text):
    for i,value in enumerate(re.split(r'`([^`]+)`',text)):
        r=p.add_run(value)
        if i%2:r.font.name='Consolas';r.font.size=Pt(10)


def section_geometry(section, start, number_format=None):
    for old in section.findall(qn('w:type')):section.remove(old)
    refs=len(section.findall(qn('w:headerReference')))+len(section.findall(qn('w:footerReference')))
    section.insert(refs,xml('type',val='nextPage'))
    mar=section.find(qn('w:pgMar'))
    for key,value in {'top':1701,'left':1440,'right':1440,'bottom':1440}.items():mar.set(qn('w:'+key),str(value))
    num=section.find(qn('w:pgNumType'))
    if num is None:num=xml('pgNumType');section.append(num)
    num.set(qn('w:start'),str(start))
    if number_format:num.set(qn('w:fmt'),number_format)
    return section


def build(reference,source,output,work):
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==REFERENCE_SHA,'template changed'
    work.mkdir(parents=True,exist_ok=True);output.parent.mkdir(parents=True,exist_ok=True)
    d=Document(reference)
    cover=[copy.deepcopy(p._p) for p in d.paragraphs[:22]]
    sections=[copy.deepcopy(s._sectPr) for s in d.sections]
    table_properties=copy.deepcopy(d.tables[0]._tbl.tblPr)
    # Source media are replaced, except the cover frame fallback retained with its vector Choice.
    for rid,rel in list(d.part.rels.items()):
        if rel.reltype.endswith('/image') and rid!='rId6':d.part.drop_rel(rid)
    b=d._element.body
    for child in list(b):b.remove(child)
    final_section=section_geometry(sections[2],1,'decimal');b.append(final_section)
    for name in ['normal','Heading 1','Heading 2','Heading 3']:
        s=d.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        for f in s.element.xpath('.//w:rFonts'):
            for a in list(f.attrib):
                if a.endswith('Theme'):del f.attrib[a]
        s.paragraph_format.widow_control=True
    normal=d.styles['normal'];normal.font.size=Pt(13)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.first_line_indent=Cm(1)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ['Heading 1','Heading 2','Heading 3']:
        s=d.styles[name];s.font.size=Pt(14);s.font.bold=True
        s.paragraph_format.line_spacing=1.0;s.paragraph_format.space_before=Pt(6);s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.left_indent=Cm(0);s.paragraph_format.first_line_indent=Cm(0)
        s.paragraph_format.keep_with_next=True
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
        for node in s.element.xpath('.//w:numPr'):node.getparent().remove(node)
    for name in ['FrontHeading','ListTitle','FigureCaption','TableCaption']:
        if name not in d.styles:d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=d.styles[name];s.base_style=normal;s.font.name='Times New Roman'
        s.font.size=Pt(14 if 'Heading' in name or name=='ListTitle' else 11)
        s.font.bold=name in ['FrontHeading','ListTitle'];s.font.italic='Caption' in name
        s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.line_spacing=1.15
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.space_after=Pt(6)
    for i in [1,2,3]:
        matches=[s for s in d.styles if (s.name or '').lower()==f'toc {i}']
        s=matches[0] if matches else d.styles.add_style(f'TOC {i}',WD_STYLE_TYPE.PARAGRAPH)
        s.font.name='Times New Roman';s.font.size=Pt(11)
        s.element.attrib.pop(qn('w:customStyle'),None)
        s.element.find(qn('w:name')).set(qn('w:val'),f'toc {i}')
        s.paragraph_format.line_spacing=1;s.paragraph_format.space_after=Pt(3)
        s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.left_indent=Cm((i-1)*.5)
    replacements={8:'BÁO CÁO ĐỒ ÁN CHUYÊN NGÀNH',9:'KỸ THUẬT PHẦN MỀM',
      11:'XÂY DỰNG ỨNG DỤNG WEB PHÁT HIỆN TRẠNG THÁI ĐEO KHẨU TRANG',
      12:'GVHD: ThS. Vũ Thị Dương',13:'Nhóm: 01',14:'Phùng Đức Anh — MSSV: Cần bổ sung',
      15:'Nguyễn Đình Luân Kiệt — MSSV: Cần bổ sung',16:'Bùi Nguyên Hải — MSSV: Cần bổ sung',
      17:'Bùi Ngọc Minh Hoàng — MSSV: Cần bổ sung',19:'Lớp: 2026IT6150003    Khóa: Cần bổ sung',21:'Hà Nội, năm 2026'}
    from docx.text.paragraph import Paragraph
    for i,node in enumerate(cover):
        if i==18:continue
        b.insert(len(b)-1,node)
        p=Paragraph(node,d._body)
        if i in replacements:
            p.text=replacements[i];p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0)
            p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name='Times New Roman';r.font.size=Pt(22 if i==11 else 16 if i in [8,9] else 13)
                r.bold=i in [8,9,11];r.font.color.rgb=RGBColor.from_string('002060' if i in [11,21] else '000000')
        if 3<=i<=7:p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    cover_section=sections[0]
    for element in list(cover_section):
        if element.tag in [qn('w:footerReference'),qn('w:pgNumType')]:cover_section.remove(element)
    def section_break(s):
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
        p._p.get_or_add_pPr().append(s)
    section_break(cover_section)
    for section in [sections[1],sections[2]]:
        for ref in section.findall(qn('w:footerReference')):
            footer=d.part.related_parts[ref.get(qn('r:id'))]._element
            for child in list(footer):footer.remove(child)
            pxml=xml('p');footer.append(pxml);p=Paragraph(pxml,d._body)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
            field(p,'PAGE','1')
    chapter='0';counts=defaultdict(int);figures=[];tables=[];front=True;first_heading=True;references=False;subheading=''
    lines=source.read_text(encoding='utf-8').splitlines();i=0
    assets=source.parent.parent/'p12/assets'
    def caption(kind,description):
        counts[(kind,chapter)]+=1;n=counts[(kind,chapter)]
        p=d.add_paragraph(style='FigureCaption' if kind=='Hình' else 'TableCaption')
        p.paragraph_format.keep_with_next=kind=='Bảng'
        p.add_run(f'{kind} {chapter}.');field(p,f'SEQ {"Hinh" if kind=="Hình" else "Bang"}{chapter}',str(n));p.add_run(' '+description)
        (figures if kind=='Hình' else tables).append({'label':f'{kind} {chapter}.{n}','description':description})
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='<!-- body -->':
            section_break(section_geometry(sections[1],1,'lowerRoman'));front=False;first_heading=True;i+=1;continue
        if line.startswith('# '):
            title=line[2:]
            references=title=='TÀI LIỆU THAM KHẢO'
            if title.startswith('CHƯƠNG '):chapter=re.search(r'CHƯƠNG (\d+)',title).group(1)
            elif title.startswith('PHỤ LỤC'):chapter=re.search(r'PHỤ LỤC ([A-Z])',title).group(1)
            style=('ListTitle' if title=='MỤC LỤC' else 'FrontHeading') if front else 'Heading 1'
            p=d.add_paragraph(title,style);p.paragraph_format.page_break_before=not first_heading;first_heading=False
            p._p.get_or_add_pPr().append(xml('numPr'));p._p.get_or_add_pPr().find(qn('w:numPr')).append(xml('numId',val=0))
        elif line.startswith('## '):
            subheading=line[3:]
            p=d.add_paragraph(line[3:],'Heading 2')
        elif line in ['{{toc}}','{{figures}}','{{tables}}']:
            instruction={'{{toc}}':'TOC \\o "1-2" \\h \\z \\t "FrontHeading,1"',
                         '{{figures}}':'TOC \\h \\z \\t "FigureCaption,1"',
                         '{{tables}}':'TOC \\h \\z \\t "TableCaption,1"'}[line]
            p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);field(p,instruction,'Cập nhật danh mục bằng Word')
        elif line.startswith('{{'):
            name=line.strip('{}');j=i+1
            while j<len(lines) and not lines[j].strip():j+=1
            description=lines[j].strip() if j<len(lines) and lines[j].startswith('Hình ') else name
            description=re.sub(r'^Hình (?:[\dA-Z]+\.\d+ )?','',description)
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
            width=11 if name in ['video','webcam'] else 15.5
            image_path=source.parent/'assets/usecases.png' if name=='usecases' else assets/f'{name}.png'
            p.add_run().add_picture(str(image_path),width=Cm(width));caption('Hình',description)
            if j<len(lines) and lines[j].startswith('Hình '):i=j
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                values=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in values):rows.append(values)
                i+=1
            if not front:
                desc={'ID':'Yêu cầu và bằng chứng kiểm chứng','Lớp test':'Precision recall và F1 theo lớp trên test',
                  'Split':'Kết quả tổng hợp validation và test' if len(rows[0])==6 else 'Phân bố bbox theo lớp và split',
                  'Thư mục và môi trường':'Lệnh tái lập theo môi trường',
                  'Truth / Prediction':'Confusion matrix trên test',
                  'Phép đo':'Kết quả HTTP và cold readiness','Luồng':'Kết quả hiển thị và độ trễ streaming',
                  'Suite P11':'Kết quả kiểm thử phần mềm P11','Evidence ID':'Danh mục evidence nội bộ'}.get(rows[0][0],rows[0][0]+' và '+rows[0][-1])
                if rows[0][0]=='Bước':desc='Luồng chính '+subheading.split(' ',1)[1]
                caption('Bảng',desc)
            t=d.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            t._tbl.replace(t._tbl.tblPr,copy.deepcopy(table_properties))
            widths=[1/len(rows[0])]*len(rows[0])
            if rows[0][0] in ['ID','Bước']:widths=[.14]+[.86/(len(rows[0])-1)]*(len(rows[0])-1)
            if rows[0][0]=='Lớp test':widths=[.32,.12,.19,.19,.18]
            if rows[0][0]=='Split' and len(rows[0])==6:widths=[.18,.18,.17,.17,.15,.15]
            for c,w in zip(t.columns,widths):c.width=Cm(15.92*w)
            for ri,row in enumerate(rows):
                cells=t.add_row().cells
                for ci,value in enumerate(row):
                    cell=cells[ci];cell.width=Cm(15.92*widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT;inline(p,value)
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=ri==0
                    props=cell._tc.get_or_add_tcPr();props.append(xml('shd',fill='E7E6E6' if ri==0 else 'FFFFFF'))
                    borders=xml('tcBorders')
                    for edge in ['top','left','bottom','right']:borders.append(xml(edge,val='single',sz=4,color='808080'))
                    props.append(borders)
                if ri==0:t.rows[0]._tr.get_or_add_trPr().append(xml('tblHeader'))
            continue
        elif line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=d.add_paragraph(lines[i]);p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
                i+=1
        else:
            p=d.add_paragraph();inline(p,line)
            if 'SHA-256:' in line or line.startswith('['):p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.first_line_indent=Cm(0)
            if references:
                p.paragraph_format.line_spacing=1.15
                for r in p.runs:r.font.size=Pt(11)
        i+=1
    settings=d.settings.element
    for child in settings.findall(qn('w:updateFields')):settings.remove(child)
    settings.append(xml('updateFields',val='true'))
    temp=work/'authored-python.docx';d.save(temp)
    editable={'word/document.xml','word/_rels/document.xml.rels','word/styles.xml','word/settings.xml','word/footer1.xml','word/footer2.xml','[Content_Types].xml'}
    with ZipFile(reference) as original,ZipFile(temp) as generated,ZipFile(output,'w',ZIP_DEFLATED) as target:
        gnames=set(generated.namelist());preserved=[]
        for name in original.namelist():
            if name.startswith('word/media/') and name!='word/media/image18.png':continue
            data=generated.read(name) if name in editable else original.read(name)
            target.writestr(name,data)
            if name not in editable:preserved.append({'part':name,'sha256':hashlib.sha256(data).hexdigest()})
        for name in gnames-set(original.namelist()):target.writestr(name,generated.read(name))
        # New images can reuse dropped source filenames; copy only those actually referenced by the new package.
        original_names=set(original.namelist())
        for name in gnames & original_names:
            if name.startswith('word/media/') and name!='word/media/image18.png':target.writestr(name,generated.read(name))
    (work/'authoring-receipt.json').write_text(json.dumps({'reference_sha256':REFERENCE_SHA,'preserved_parts':preserved,'figures':figures,'tables':tables},ensure_ascii=False,indent=2),encoding='utf-8')
    (work/'authored-package.docx').write_bytes(output.read_bytes())
    print(f'Authored {len(figures)} figures and {len(tables)} captioned tables from retained reference.')


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--reference',type=Path,required=True);p.add_argument('--output',type=Path,required=True);p.add_argument('--work',type=Path,required=True)
    args=p.parse_args();build(args.reference,Path(__file__).with_name('bao-cao.md'),args.output,args.work)
